import argparse
import glob
import io
import os
import re
import sys
from collections import Counter, defaultdict

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

TACH_CAU = re.compile(r"(?<=[.;:!?])\s+")
KHONG_CHU = re.compile(r"[^0-9A-Za-zÀ-ỹ ]+")


def chuan_hoa(van_ban):
    return KHONG_CHU.sub(" ", van_ban.lower()).split()


def doc_docx(duong_dan):
    import docx
    tai_lieu = docx.Document(duong_dan)
    phan = [p.text for p in tai_lieu.paragraphs]
    for bang in tai_lieu.tables:
        for hang in bang.rows:
            phan.extend(o.text for o in hang.cells)
    return "\n".join(t.strip() for t in phan if t and t.strip())


def doc_pdf(duong_dan):
    import pymupdf
    with pymupdf.open(duong_dan) as tai_lieu:
        return "".join(trang.get_text() for trang in tai_lieu)


def ngram(tu, n):
    return [tuple(tu[i:i + n]) for i in range(len(tu) - n + 1)]


def cau_lap_noi_bo(van_ban, so_tu_toi_thieu):
    dem = Counter()
    vi_tri = defaultdict(list)
    for dong in van_ban.split("\n"):
        for cau in TACH_CAU.split(dong):
            tu = chuan_hoa(cau)
            if len(tu) < so_tu_toi_thieu:
                continue
            khoa = " ".join(tu)
            dem[khoa] += 1
            vi_tri[khoa].append(cau.strip()[:110])
    return [(k, v, vi_tri[k][0]) for k, v in dem.most_common() if v > 1]


def trung_voi_nguon(van_ban_bao_cao, cac_nguon, do_dai_ngram):
    tu_bao_cao = chuan_hoa(van_ban_bao_cao)
    tap_bao_cao = set(ngram(tu_bao_cao, do_dai_ngram))
    ket_qua = {}
    for ten, noi_dung in cac_nguon.items():
        tap_nguon = set(ngram(chuan_hoa(noi_dung), do_dai_ngram))
        chung = tap_bao_cao & tap_nguon
        ket_qua[ten] = sorted(" ".join(g) for g in chung)
    return ket_qua, len(tap_bao_cao)


def main():
    parser = argparse.ArgumentParser(
        description="Đo trùng lặp nội bộ của báo cáo và trùng lặp nguyên văn với ba bài báo nguồn")
    parser.add_argument("--bao_cao", default="bao_cao/output/bao_cao_bonevqa_prompt.docx")
    parser.add_argument("--thu_muc_nguon", default="yeucau")
    parser.add_argument("--so_tu_cau", type=int, default=12,
                        help="Số từ tối thiểu để coi một câu là đáng kể khi đếm lặp nội bộ")
    parser.add_argument("--ngram", type=int, default=8,
                        help="Độ dài chuỗi từ liên tiếp để coi là trùng nguyên văn với nguồn")
    parser.add_argument("--hien", type=int, default=15)
    args = parser.parse_args()

    van_ban = doc_docx(args.bao_cao)
    tong_tu = len(chuan_hoa(van_ban))
    print(f"Báo cáo: {len(van_ban)} ký tự, {tong_tu} từ\n")

    lap = cau_lap_noi_bo(van_ban, args.so_tu_cau)
    tong_lan_thua = sum(v - 1 for _, v, _ in lap)
    print(f"== TRÙNG LẶP NỘI BỘ (câu từ {args.so_tu_cau} từ trở lên, xuất hiện nhiều hơn một lần) ==")
    print(f"Số câu bị lặp: {len(lap)} | Tổng số lần thừa: {tong_lan_thua}")
    for khoa, so_lan, vi_du in lap[:args.hien]:
        print(f"  {so_lan} lần · {len(khoa.split())} từ · {vi_du}")
    print()

    nguon = {}
    for duong_dan in sorted(glob.glob(os.path.join(args.thu_muc_nguon, "*.pdf"))):
        nguon[os.path.basename(duong_dan)] = doc_pdf(duong_dan)
    ket_qua, so_ngram = trung_voi_nguon(van_ban, nguon, args.ngram)
    print(f"== TRÙNG NGUYÊN VĂN VỚI BÀI BÁO NGUỒN (chuỗi {args.ngram} từ liên tiếp) ==")
    print(f"Báo cáo có {so_ngram} chuỗi {args.ngram} từ khác nhau")
    for ten, chuoi in ket_qua.items():
        ty_le = len(chuoi) / max(so_ngram, 1) * 100
        print(f"  {ten[:52]:54s} {len(chuoi):4d} chuỗi trùng ({ty_le:.3f}%)")
        for c in chuoi[:args.hien]:
            print(f"      · {c}")
    print()


if __name__ == "__main__":
    main()
